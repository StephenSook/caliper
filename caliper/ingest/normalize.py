"""Normalize the three supplied QA exports into one long table.

The three exports are not three datasets. They are three instruments pointed at
the same seventeen evaluation events, which is what makes cross domain
corroboration on the same call possible at all.

Columns are resolved by HEADER NAME, not by letter. The Compliance export
carries an extra "Date of Evaluation" column that shifts every later column by
one, and a positional map silently reads the wrong field.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .pii_scan import scan_many
from .pseudonymize import new_salt, pseudonymize

DOMAIN_FILES = {
    "member_experience": "QA Raw Data - Member Experience Focus (Healthcare Partner).xlsx",
    "business_process": "Call Flow - Business Process (Healthcare Partner).xlsx",
    "compliance": "Compliance Raw Data (Healthcare Partner).xlsx",
}

DOMAIN_PREFIX = {
    "member_experience": "ME",
    "business_process": "BP",
    "compliance": "CP",
}

# Header name -> canonical field. Resolved case insensitively and whitespace
# collapsed, because supplied exports are hand maintained.
HEADER_MAP = {
    "qa name": "rater",
    "agentnames": "agent",
    "agent names": "agent",
    "date of call": "call_date",
    "questions": "item_text",
    "answers": "answer",
}


@dataclass(frozen=True)
class Observation:
    """One scored item on one evaluation event.

    `call_date` is kept only to order occasions and to group items belonging to
    the same evaluation. It is deliberately EXCLUDED from `as_dict`, because our
    own identifier scanner treats a date of service as an identifier, and it
    would be incoherent to scan supplied material for dates and then render the
    supplied dates ourselves. Anything downstream sees `occasion_index` instead.
    """

    eval_id: str
    agent_ref: str
    rater_ref: str
    domain: str
    item_id: str
    item_text: str
    passed: bool | None
    call_date: str = field(repr=False, compare=True)
    occasion_index: int = 0

    def as_dict(self) -> dict:
        """Serialization boundary. No supplied identifier crosses it."""
        data = asdict(self)
        data.pop("call_date", None)
        return data


def _norm_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def slugify_item(domain: str, text: str, taken: set[str]) -> str:
    """Stable, readable item id, e.g. ME-EXPLAINED-OPTIONS-RESOLUTION."""
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    words = re.findall(r"[A-Za-z]+", ascii_text)
    stop = {"the", "a", "an", "and", "or", "to", "of", "in", "all", "any", "for", "was", "did"}
    keep = [w.upper() for w in words if w.lower() not in stop][:4]
    base = f"{DOMAIN_PREFIX[domain]}-" + "-".join(keep or ["ITEM"])
    candidate, n = base, 2
    while candidate in taken:
        candidate = f"{base}-{n}"
        n += 1
    taken.add(candidate)
    return candidate


def _cell_date(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value or "").strip()


def _parse_answer(value) -> bool | None:
    token = str(value or "").strip().upper()
    if token in {"YES", "Y", "PASS", "1", "TRUE"}:
        return True
    if token in {"NO", "N", "FAIL", "0", "FALSE"}:
        return False
    return None


def load_domain(path: Path, domain: str, salt: str) -> tuple[list[Observation], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)

    header = next(rows)
    index: dict[str, int] = {}
    for i, cell in enumerate(header):
        field = HEADER_MAP.get(_norm_header(cell))
        if field and field not in index:
            index[field] = i
    missing = {"rater", "agent", "call_date", "item_text", "answer"} - index.keys()
    if missing:
        raise ValueError(f"{path.name}: could not resolve columns {sorted(missing)}")

    observations: list[Observation] = []
    slugs: dict[str, str] = {}
    taken: set[str] = set()
    scanned: list[object] = []

    for row in rows:
        if row is None:
            continue
        text = str(row[index["item_text"]] or "").strip()
        agent = str(row[index["agent"]] or "").strip()
        if not text or not agent:
            continue

        scanned.extend(row)
        if text not in slugs:
            slugs[text] = slugify_item(domain, text, taken)

        agent_ref = pseudonymize(agent, salt, "A")
        call_date = _cell_date(row[index["call_date"]])
        observations.append(
            Observation(
                eval_id=pseudonymize(f"{agent}|{call_date}", salt, "E"),
                agent_ref=agent_ref,
                rater_ref=pseudonymize(str(row[index["rater"]] or "").strip(), salt, "R"),
                domain=domain,
                item_id=slugs[text],
                item_text=text,
                passed=_parse_answer(row[index["answer"]]),
                call_date=call_date,
            )
        )

    workbook.close()
    return _assign_occasions(observations), scan_many(scanned).counts


def _assign_occasions(observations: list[Observation]) -> list[Observation]:
    """Replace a real date with that agent's 1-based occasion ordinal.

    Ordering still uses the true date so "first call" means first, but the date
    itself never reaches a rendered surface.
    """
    order: dict[str, list[str]] = {}
    for o in observations:
        dates = order.setdefault(o.agent_ref, [])
        if o.call_date not in dates:
            dates.append(o.call_date)
    for dates in order.values():
        dates.sort()
    return [
        Observation(
            eval_id=o.eval_id,
            agent_ref=o.agent_ref,
            rater_ref=o.rater_ref,
            domain=o.domain,
            item_id=o.item_id,
            item_text=o.item_text,
            passed=o.passed,
            call_date=o.call_date,
            occasion_index=order[o.agent_ref].index(o.call_date) + 1,
        )
        for o in observations
    ]


def load_all(data_dir: str | Path, salt: str | None = None):
    """Load all three domains. Returns (observations, redaction counts, salt)."""
    directory = Path(data_dir)
    salt = salt or new_salt()
    observations: list[Observation] = []
    redactions: dict[str, int] = {}
    for domain, filename in DOMAIN_FILES.items():
        path = directory / filename
        if not path.exists():
            raise FileNotFoundError(f"missing supplied export: {path}")
        rows, counts = load_domain(path, domain, salt)
        observations.extend(rows)
        for kind, count in counts.items():
            redactions[kind] = redactions.get(kind, 0) + count
    return observations, redactions, salt
