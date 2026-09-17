"""The sponsor's six tab design workbook, emitted from one run.

Their workbook completes six tabs in order, each carrying a purple box telling a
designer to copy a prompt, fill in the brackets, attach files and send. That is
six manual re prompts across six tabs, and the rubric line worth twenty five
percent is called "Seamless End-to-End Generation (No Reprompting)".

These tests hold the emitted file to the sponsor's own schema, because a workbook
that is merely "inspired by" theirs does not open in the tool they already use.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from caliper.export.rcx_workbook import TABS, to_bytes

pytest.importorskip("openpyxl")

DIAGNOSIS = {
    "diagnosis_id": "DIAG-001",
    "domain": "member_experience",
    "behavior": "Explained options and resolution clearly",
    "observed": {
        "fails": 14,
        "denominator": 17,
        "rate": 0.8235,
        "breadth_subjects": 8,
        "breadth_denominator": 10,
        "difficulty_p": 0.176,
        "discrimination_rpb": 0.083,
        "item_flags": [],
    },
    "existing_curriculum_coverage": {"covered": True, "source": "Week 6"},
    "root_cause_primary": "MEASUREMENT_STANDARD_SETTING",
    "recommended_intervention_class": "Rewrite the item and calibrate the evaluators.",
    "alternatives_rejected": [{"cause": "SKILL", "reason": "already covered"}],
    "individual_attribution": {"statement": "Declined, with the remedy attached."},
}

AUDIT = {
    "domains": {
        "member_experience": {
            "n_evaluations": 17,
            "n_items": 9,
            "reliability": {
                "statistic": "KR-20",
                "point_estimate": 0.4661,
                "ci_low": -0.0192,
                "ci_high": 0.7756,
                "verdict_reason": "The sample cannot establish whether it works.",
            },
            "summary": {"items_outside_difficulty_band": 5, "items_with_zero_variance": 2},
        }
    },
    "connectivity": {"linkage_fragility": 2, "n_subjects": 10, "calls_double_scored": 0, "n_calls": 17},
}

BUNDLE = {
    "rewritten_item": {"item_id": "ME-X-v2-CONFIRM-UNDERSTANDING"},
    "objectives": [
        {
            "objective_id": "OBJ-001",
            "behavior": "elicit a restatement of the amount owed",
            "criterion": "restatement substantively correct on the first attempt",
        }
    ],
    "activities": [
        {
            "activity_id": "ACT-001",
            "title": "Say it back",
            "activity_type": "Role play",
            "duration_minutes": 12,
        }
    ],
    "simulations": [{"simulation_id": "SIM-001", "call_driver": "EOB dispute"}],
    "metrics": [
        {"metric_id": "METRIC-001", "baseline_rate": 0.176, "target_rate": 0.60, "required_n_per_arm": 20}
    ],
    "knowledge_checks": [
        {
            "kc_id": "KC-001",
            "question_stem": "What is the best next step?",
            "options": [
                "Ask her to restate it",
                "Read the lines in order",
                "Apply a credit now",
                "Transfer to billing",
            ],
            "correct_option": "A",
            "sme_confirmed": True,
        }
    ],
    "outcome_claim": {"verdict": "INSUFFICIENT_POWER", "statement": "We will not claim improvement."},
    "alignment": {"kc_gate": {"passed": 1, "total": 1}},
}

PERSONA = {
    "disclosure": "Synthetic member over real public plan mechanics.",
    "member_state": {
        "emotional_register": "anxious",
        "believes": "I owe nothing",
        "core_misunderstanding": "thinks coinsurance is a copay",
        "opening_line": "I already met my deductible.",
    },
    "plan_facts": {
        "billed_amount": 1700.0,
        "allowed_amount": 1425.0,
        "network_adjustment": 275.0,
        "plan_paid": 1140.0,
        "member_responsibility": 285.0,
        "coinsurance_rate": 0.20,
    },
    "scored_criteria": [
        {"id": "VERIFY_IDENTITY", "text": "Identity verified first", "basis": "45 CFR 164.514(h)"},
        {"id": "CONFIRM_UNDERSTANDING", "text": "Member restated what she owes"},
    ],
}


def build() -> object:
    data = to_bytes(run_id="RUN-TEST01", audit=AUDIT, diagnosis=DIAGNOSIS, bundle=BUNDLE, persona=PERSONA)
    return load_workbook(BytesIO(data))


def all_text(ws) -> str:
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_the_six_tabs_match_the_sponsor_workbook_exactly_and_in_order():
    """Not "inspired by" theirs. The same tab names, in the same sequence, so the
    output opens in the tool a designer already uses."""
    assert build().sheetnames == TABS


def test_every_tab_is_populated():
    """A workbook with the right tab names and empty cells is worse than none."""
    wb = build()
    for name in TABS:
        filled = sum(1 for row in wb[name].iter_rows() for c in row if c.value not in (None, ""))
        assert filled > 5, f"{name} is effectively empty with {filled} cells"


def test_every_tab_carries_the_run_and_diagnosis_it_came_from():
    wb = build()
    for name in TABS:
        text = all_text(wb[name])
        assert "RUN-TEST01" in text, f"{name} does not say which run produced it"
        assert "DIAG-001" in text, f"{name} does not trace to a diagnosis"


def test_the_defect_table_carries_the_measured_figures():
    text = all_text(build()["QA Needs Analysis"])
    assert "0.176" in text and "0.083" in text
    assert "14 of 17" in text
    assert "MEASUREMENT" in text.upper()


def test_the_simulation_is_scored_on_the_rewritten_item():
    """The check that matters most, carried into the deliverable: a simulation
    scored on any other item measures the wrong behaviour."""
    text = all_text(build()["QA Needs Analysis"])
    assert BUNDLE["rewritten_item"]["item_id"] in text


def test_unsupplied_data_is_marked_unsupplied_rather_than_invented():
    """Their template has CSAT and QA trend tables we were given no data for.
    Filling them with plausible numbers is the fabrication the case penalises."""
    text = all_text(build()["QA Needs Analysis"])
    assert "not supplied" in text.lower()
    assert "do not fabricate" in text.lower() or "rather than estimated" in text.lower()


def test_the_persona_is_labelled_synthetic_in_the_deliverable():
    text = all_text(build()["Persona Details"])
    assert "synthetic" in text.lower()
    assert "unnamed by design" in text.lower()


def test_the_persona_tab_carries_consistent_plan_arithmetic():
    text = all_text(build()["Persona Details"])
    for value in ("1700", "1425", "285", "1140"):
        assert value in text, f"the deliverable dropped {value} from the plan facts"


def test_identity_verification_is_a_scored_skill_in_the_skills_tab():
    text = all_text(build()["Simulation Skills Outline"])
    assert "verify identity" in text.lower()
    assert "164.514" in text


def test_the_knowledge_check_tab_names_the_sponsors_own_five_check_gate():
    text = all_text(build()["Knowledge Check + Activity"]).lower()
    for phrase in ("objective aligned", "decision based", "plausible", "sop aligned", "sme"):
        assert phrase in text, f"the gate description is missing {phrase}"
